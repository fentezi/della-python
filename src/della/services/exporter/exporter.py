"""Excel exporter for cargo cards."""

from __future__ import annotations

import os
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font

from della.services.exporter.formatters import (
    calculate_max_cities,
    create_headers,
    format_location,
    format_price,
    format_tags,
    generate_file_name,
    get_string_value,
    split_date_range,
)
from della.services.parser.models import CargoCard


class Exporter:
    """Excel exporter for cargo cards."""

    def __init__(self, output_dir: str = "./exports", filename_prefix: str = "cargo_export"):
        """Initialize exporter.

        Args:
            output_dir: Directory for exported files.
            filename_prefix: Prefix for exported filenames.
        """
        self.output_dir = output_dir
        self.filename_prefix = filename_prefix

    def export_to_excel(self, cards: List[CargoCard]) -> str:
        """Export cards to Excel file.

        Args:
            cards: List of cargo cards to export.

        Returns:
            Path to created Excel file.
        """
        # Create directory if doesn't exist
        os.makedirs(self.output_dir, exist_ok=True)

        # Create new Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Картки"

        # If no cards, create file with headers only
        if not cards:
            headers = create_headers(0, 0)
            self._write_headers(ws, headers)
        else:
            # Determine maximum number of cities
            max_from, max_to = calculate_max_cities(cards)

            # Create and write headers
            headers = create_headers(max_from, max_to)
            self._write_headers(ws, headers)

            # Write card data
            self._write_card_data(ws, cards, max_from, max_to)

        # Generate filename and full path
        filename = generate_file_name(self.filename_prefix)
        filepath = Path(self.output_dir) / filename

        # Save file
        wb.save(str(filepath))

        return str(filepath)

    def _write_headers(self, ws, headers: List[str]) -> None:
        """Write headers to Excel worksheet.

        Args:
            ws: Worksheet object.
            headers: List of header strings.
        """
        # Create bold font style for headers
        bold_font = Font(bold=True)

        # Write headers to first row
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = bold_font

    def _write_card_data(
        self,
        ws,
        cards: List[CargoCard],
        max_from: int,
        max_to: int,
    ) -> None:
        """Write card data to Excel worksheet.

        Args:
            ws: Worksheet object.
            cards: List of cargo cards.
            max_from: Maximum 'From' cities count.
            max_to: Maximum 'To' cities count.
        """
        for row_idx, card in enumerate(cards, start=2):  # Start from row 2 (1 is headers)
            col_idx = 1

            # Split date into start and end
            start_date, end_date = split_date_range(card.date)

            # Fixed fields
            values = [
                card.request_id,
                start_date,
                end_date,
                get_string_value(card.weight),
                get_string_value(card.volume),
                card.truck_type,
                card.cargo_type,
            ]

            # Write fixed fields
            for value in values:
                ws.cell(row=row_idx, column=col_idx, value=value)
                col_idx += 1

            # Write FromCities (dynamic number of columns)
            for i in range(max_from):
                value = ""
                if i < len(card.from_cities):
                    value = format_location(card.from_cities[i])
                ws.cell(row=row_idx, column=col_idx, value=value)
                col_idx += 1

            # Write ToCities (dynamic number of columns)
            for i in range(max_to):
                value = ""
                if i < len(card.to_cities):
                    value = format_location(card.to_cities[i])
                ws.cell(row=row_idx, column=col_idx, value=value)
                col_idx += 1

            # Write remaining fixed fields
            main_price, price_per_km, price_tags = format_price(card.price)
            final_values = [
                format_tags(card.request_tags),
                main_price,
                price_per_km,
                price_tags,
            ]

            for value in final_values:
                ws.cell(row=row_idx, column=col_idx, value=value)
                col_idx += 1
